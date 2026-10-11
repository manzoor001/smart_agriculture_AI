from pathlib import Path
from pypdf import PdfReader
from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
RAW_DIR = BASE_DIR / "raw"
OUTPUT_DIR = BASE_DIR / "processed"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def extract_pdf(path):
    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            pages.append(
                f"\n--- Page {page_number} ---\n{text}"
            )

    return "\n".join(pages)


def extract_docx(path):
    doc = Document(str(path))
    paragraphs = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.text.strip()
    ]

    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(
                " | ".join(cell.text for cell in row.cells)
            )

    return "\n".join(paragraphs)


def main():
    files = (
        list(RAW_DIR.glob("*.pdf"))
        + list(RAW_DIR.glob("*.docx"))
    )

    if not files:
        print(f"No PDF or DOCX files found in {RAW_DIR}")
        return

    for path in files:
        try:
            print(f"\nExtracting: {path.name}")

            if path.suffix.lower() == ".pdf":
                text = extract_pdf(path)
            else:
                text = extract_docx(path)

            if not text.strip():
                print("No text extracted. This may be a scanned document.")
                continue

            output_path = OUTPUT_DIR / f"{path.stem}.txt"
            output_path.write_text(text, encoding="utf-8")

            print(f"Saved: {output_path}")
            print(f"Characters extracted: {len(text)}")

        except Exception as error:
            print(f"Failed: {path.name}: {error}")


if __name__ == "__main__":
    main()
